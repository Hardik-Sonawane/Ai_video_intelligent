"""
Script to generate the comprehensive Academic Research Paper and Technical Project Report
for the 'AI Video Intelligence' project in Word (.docx) format.
"""

import os
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Colors
COLOR_PRIMARY_NAVY = RGBColor(15, 41, 66)      # #0F2942
COLOR_SECONDARY_BLUE = RGBColor(43, 108, 176)   # #2B6CB0
COLOR_TEAL_ACCENT = RGBColor(49, 151, 149)      # #319795
COLOR_BODY_TEXT = RGBColor(45, 55, 72)          # #2D3748
COLOR_MUTED_GRAY = RGBColor(113, 128, 150)      # #718096
HEX_LIGHT_BG = "F7FAFC"
HEX_BORDER_GRAY = "CBD5E0"
HEX_PRIMARY_NAVY = "0F2942"
HEX_SECONDARY_BLUE = "2B6CB0"
HEX_CALLOUT_BG = "EDF2F7"
HEX_CODE_BG = "F8F9FA"


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell margins in twentieths of a point (dxa)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_background(cell, hex_color):
    """Set the background color of a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    """
    Set cell borders:
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='FF0000')
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
    tcPr.append(tcBorders)


def add_document_header_footer(doc):
    """Configure running header and footer."""
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("AI Video Intelligence: Agentic Multimodal Architecture  |  Technical Research Report")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED_GRAY

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("AI Video Intelligence  •  Advanced AI & LLM Systems  •  Confidential / Academic Use")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = COLOR_MUTED_GRAY


def add_title(doc, title_text, subtitle_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title_text)
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY_NAVY

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(16)
    run2 = p2.add_run(subtitle_text)
    run2.font.name = "Calibri"
    run2.font.size = Pt(13)
    run2.font.bold = True
    run2.font.color.rgb = COLOR_SECONDARY_BLUE


def add_author_block(doc, author_name, affiliation, date_str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(20)

    r1 = p.add_run(f"{author_name}\n")
    r1.font.name = "Calibri"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_BODY_TEXT

    r2 = p.add_run(f"{affiliation}\n")
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    r2.font.italic = True
    r2.font.color.rgb = COLOR_MUTED_GRAY

    r3 = p.add_run(f"Project Repository: project_ai_iv  |  Evaluation Date: {date_str}")
    r3.font.name = "Calibri"
    r3.font.size = Pt(9.5)
    r3.font.color.rgb = COLOR_MUTED_GRAY


def add_abstract(doc, abstract_text, keywords_list):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F4F6F9")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    set_cell_border(cell,
                    top=dict(val='single', sz=6, color=HEX_SECONDARY_BLUE),
                    bottom=dict(val='single', sz=6, color=HEX_SECONDARY_BLUE),
                    left=dict(val='single', sz=18, color=HEX_PRIMARY_NAVY),
                    right=dict(val='single', sz=6, color=HEX_SECONDARY_BLUE))

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    r_title = p.add_run("ABSTRACT\n")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(11)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY_NAVY

    r_body = p.add_run(abstract_text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = COLOR_BODY_TEXT

    p_kw = cell.add_paragraph()
    p_kw.paragraph_format.space_before = Pt(6)
    p_kw.paragraph_format.space_after = Pt(0)
    r_kw_title = p_kw.add_run("Keywords: ")
    r_kw_title.font.name = "Calibri"
    r_kw_title.font.size = Pt(9.5)
    r_kw_title.font.bold = True
    r_kw_title.font.color.rgb = COLOR_PRIMARY_NAVY

    r_kw_body = p_kw.add_run("; ".join(keywords_list))
    r_kw_body.font.name = "Calibri"
    r_kw_body.font.size = Pt(9.5)
    r_kw_body.font.italic = True
    r_kw_body.font.color.rgb = COLOR_BODY_TEXT

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(12)


def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY_NAVY
    return p


def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY_BLUE
    return p


def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_BODY_TEXT
    return p


def add_paragraph(doc, text, bold_prefix="", italic=False, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = align

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BODY_TEXT

    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.italic = italic
    r.font.color.rgb = COLOR_BODY_TEXT
    return p


def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_BODY_TEXT

    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_BODY_TEXT
    return p


def add_callout(doc, text, title="KEY ARCHITECTURAL HIGHLIGHT", border_color=HEX_SECONDARY_BLUE, bg_color=HEX_CALLOUT_BG):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    set_cell_border(cell,
                    top=dict(val='none'),
                    bottom=dict(val='none'),
                    left=dict(val='single', sz=24, color=border_color),
                    right=dict(val='none'))

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r_t = p.add_run(f"📌 {title}\n")
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(10)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_PRIMARY_NAVY

    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(9.5)
    r.font.italic = True
    r.font.color.rgb = COLOR_BODY_TEXT

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)


def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, HEX_CODE_BG)
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    set_cell_border(cell,
                    top=dict(val='single', sz=4, color=HEX_BORDER_GRAY),
                    bottom=dict(val='single', sz=4, color=HEX_BORDER_GRAY),
                    left=dict(val='single', sz=12, color=HEX_PRIMARY_NAVY),
                    right=dict(val='single', sz=4, color=HEX_BORDER_GRAY))

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)


def add_formula_block(doc, formula_text, formula_id):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c1 = table.cell(0, 0)
    c2 = table.cell(0, 1)

    c1.width = Inches(5.8)
    c2.width = Inches(0.7)

    set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
    set_cell_margins(c2, top=60, bottom=60, left=50, right=50)

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(formula_text)
    r1.font.name = "Cambria Math"
    r1.font.size = Pt(10.5)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_PRIMARY_NAVY

    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run(f"({formula_id})")
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    r2.font.color.rgb = COLOR_MUTED_GRAY


def add_styled_table(doc, headers, data, col_widths=None, alignments=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Format header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], HEX_PRIMARY_NAVY)
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = alignments[i] if alignments else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    # Format data rows
    for r_idx, row in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "FFFFFF" if r_idx % 2 == 0 else "F7FAFC"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=100, right=100)
            set_cell_border(row_cells[c_idx],
                            top=dict(val='single', sz=2, color=HEX_BORDER_GRAY),
                            bottom=dict(val='single', sz=2, color=HEX_BORDER_GRAY),
                            left=dict(val='single', sz=2, color=HEX_BORDER_GRAY),
                            right=dict(val='single', sz=2, color=HEX_BORDER_GRAY))
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = alignments[c_idx] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9.0)
                run.font.color.rgb = COLOR_BODY_TEXT

    # Apply widths
    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)


print("Helper functions defined successfully.")
